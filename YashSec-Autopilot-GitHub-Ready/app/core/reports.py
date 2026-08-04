from __future__ import annotations

import html
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.config import REPORT_DIR


SEVERITY_RANK = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}


def _sorted_findings(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(findings, key=lambda item: (SEVERITY_RANK.get(item.get("severity", "medium"), 9), item.get("title", "")))


def generate_json(scan: dict[str, Any], repo: dict[str, Any], findings: list[dict[str, Any]]) -> Path:
    path = REPORT_DIR / f"scan-{scan['id']}-report.json"
    path.write_text(
        json.dumps({"repository": repo, "scan": scan, "findings": findings}, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    return path


def generate_html(scan: dict[str, Any], repo: dict[str, Any], findings: list[dict[str, Any]]) -> Path:
    path = REPORT_DIR / f"scan-{scan['id']}-report.html"
    summary = scan.get("summary") or {}
    finding_rows = []
    for finding in _sorted_findings(findings):
        finding_rows.append(
            f"""
            <tr>
              <td><span class="sev {html.escape(finding['severity'])}">{html.escape(finding['severity'].upper())}</span></td>
              <td><strong>{html.escape(finding['title'])}</strong><br><small>{html.escape(finding.get('description') or '')}</small></td>
              <td>{html.escape(finding.get('tool') or '')}</td>
              <td>{html.escape(finding.get('file_path') or '—')}:{html.escape(str(finding.get('line_start') or '—'))}</td>
              <td>{html.escape(finding.get('review_status') or 'open')}</td>
            </tr>
            """
        )
    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>YashSec Report</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;margin:0;background:#f4f6f9;color:#18202b}}
.header{{background:#111827;color:white;padding:36px 52px}} .header h1{{margin:0;font-size:30px}}
.header p{{color:#cbd5e1;margin:8px 0 0}} .wrap{{padding:34px 52px}}
.cards{{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;margin:20px 0}}
.card{{background:white;border:1px solid #e5e7eb;border-radius:12px;padding:16px}}
.card b{{display:block;font-size:24px;margin-top:8px}} table{{width:100%;border-collapse:collapse;background:white;border-radius:12px;overflow:hidden}}
th,td{{padding:12px;text-align:left;border-bottom:1px solid #e5e7eb;vertical-align:top}} th{{background:#f8fafc}}
.sev{{font-size:11px;font-weight:700;padding:5px 8px;border-radius:99px}} .critical{{background:#fee2e2;color:#991b1b}} .high{{background:#ffedd5;color:#9a3412}} .medium{{background:#fef3c7;color:#92400e}} .low{{background:#dbeafe;color:#1e40af}} .info{{background:#e5e7eb;color:#374151}}
small{{color:#64748b}} .meta{{display:grid;grid-template-columns:1fr 1fr;gap:10px;background:white;border:1px solid #e5e7eb;border-radius:12px;padding:18px}}
@media print{{body{{background:white}} .wrap{{padding:18px}}}}
</style></head><body>
<div class="header"><h1>YashSec Autopilot Security Report</h1><p>{html.escape(repo['name'])} · Scan #{scan['id']} · {html.escape(scan.get('completed_at') or scan.get('created_at') or '')}</p></div>
<div class="wrap">
<div class="meta"><div><b>Repository path</b><br>{html.escape(repo.get('local_path') or '')}</div><div><b>Coverage</b><br>{html.escape(', '.join(summary.get('completed_tools') or []))}</div></div>
<div class="cards">
<div class="card">Critical<b>{summary.get('critical',0)}</b></div><div class="card">High<b>{summary.get('high',0)}</b></div><div class="card">Medium<b>{summary.get('medium',0)}</b></div><div class="card">Low<b>{summary.get('low',0)}</b></div><div class="card">Total<b>{summary.get('total',len(findings))}</b></div>
</div>
<h2>Findings</h2><table><thead><tr><th>Severity</th><th>Finding</th><th>Tool</th><th>Location</th><th>Status</th></tr></thead><tbody>{''.join(finding_rows) or '<tr><td colspan="5">No findings recorded.</td></tr>'}</tbody></table>
</div></body></html>"""
    path.write_text(document, encoding="utf-8")
    return path


def _set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9)


def generate_docx(scan: dict[str, Any], repo: dict[str, Any], findings: list[dict[str, Any]]) -> Path:
    path = REPORT_DIR / f"scan-{scan['id']}-report.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(13)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("YashSec Autopilot\n")
    subtitle = title.add_run("Security Review Report")
    subtitle.font.size = Pt(16)
    subtitle.font.bold = False

    document.add_paragraph(f"Repository: {repo['name']}")
    document.add_paragraph(f"Scan ID: {scan['id']}  |  Completed: {scan.get('completed_at') or 'Not completed'}")

    summary = scan.get("summary") or {}
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        f"The automated review recorded {summary.get('total', len(findings))} finding(s): "
        f"{summary.get('critical', 0)} critical, {summary.get('high', 0)} high, "
        f"{summary.get('medium', 0)} medium, {summary.get('low', 0)} low, and {summary.get('info', 0)} informational. "
        "Automated findings require human validation before production decisions are made."
    )

    coverage = document.add_table(rows=1, cols=3)
    coverage.alignment = WD_TABLE_ALIGNMENT.LEFT
    coverage.style = "Light Shading Accent 1"
    for index, value in enumerate(("Completed tools", "Skipped or missing", "API target")):
        _set_cell_text(coverage.rows[0].cells[index], value, bold=True)
    _set_cell_text(coverage.add_row().cells[0], ", ".join(summary.get("completed_tools") or []) or "None")
    _set_cell_text(coverage.rows[1].cells[1], ", ".join(summary.get("skipped_tools") or []) or "None")
    _set_cell_text(coverage.rows[1].cells[2], scan.get("api_url") or "Not supplied/discovered")

    document.add_heading("Findings", level=1)
    for index, finding in enumerate(_sorted_findings(findings), start=1):
        document.add_heading(f"{index}. {finding['title']}", level=2)
        meta = document.add_table(rows=2, cols=4)
        meta.style = "Light Grid Accent 1"
        labels = ("Severity", "Tool", "Status", "Location")
        values = (
            finding.get("severity", "").upper(),
            finding.get("tool", ""),
            finding.get("review_status", "open"),
            f"{finding.get('file_path') or '—'}:{finding.get('line_start') or '—'}",
        )
        for cell, label in zip(meta.rows[0].cells, labels):
            _set_cell_text(cell, label, bold=True)
        for cell, value in zip(meta.rows[1].cells, values):
            _set_cell_text(cell, str(value))
        document.add_paragraph(finding.get("description") or "")
        if finding.get("evidence"):
            paragraph = document.add_paragraph()
            paragraph.add_run("Evidence\n").bold = True
            paragraph.add_run(str(finding["evidence"])[:3500])
        paragraph = document.add_paragraph()
        paragraph.add_run("Recommended remediation\n").bold = True
        paragraph.add_run(finding.get("remediation") or "Review and remediate the affected component.")
        if finding.get("reviewer_notes"):
            paragraph = document.add_paragraph()
            paragraph.add_run("Reviewer notes\n").bold = True
            paragraph.add_run(finding["reviewer_notes"])

    document.add_heading("Limitations", level=1)
    document.add_paragraph(
        "This report combines configured automated scanners and safe local checks. Missing tools, unavailable runtime targets, "
        "authentication context, unsupported languages, or absent OpenAPI documentation reduce coverage. The absence of a finding "
        "does not prove the absence of a vulnerability."
    )
    document.save(path)
    return path
